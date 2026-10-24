# -*- coding: utf-8 -*-
import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_document():
    doc = docx.Document()
    
    # Page setup: Letter, 0.75 in margins
    for sec in doc.sections:
        sec.top_margin = Inches(0.75)
        sec.bottom_margin = Inches(0.75)
        sec.left_margin = Inches(0.75)
        sec.right_margin = Inches(0.75)
        sec.page_width = Inches(8.5)
        sec.page_height = Inches(11.0)
        
    black = RGBColor(0, 0, 0)
    
    # Base Style
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = black

    views_data = [
        {
            'num': 1,
            'file': 'Screenshot 2026-10-20 090245.png',
            'title': 'Pabellón de Fiestas y Ceremonias (Detalle de Exhibición)',
            'desc': 'Vista frontal de exhibición del Pabellón de Fiestas con cuadros de estambre e instrumentos.'
        },
        {
            'num': 2,
            'file': 'Screenshot 2026-10-20 101123.png',
            'title': 'Pabellón de Fiestas y Ceremonias (Cédulas en Wixárika)',
            'desc': 'Vista frontal de Fiestas y Ceremonias con cédulas en wixárika y tambor ceremonial Tepu.'
        },
        {
            'num': 3,
            'file': 'Screenshot 2026-10-20 122018.png',
            'title': 'Pabellón de Gastronomía Tradicional (Montaje de Mesas)',
            'desc': 'Vista frontal de Gastronomía mostrando la doble mesa de platillos y repisa de elotes.'
        },
        {
            'num': 4,
            'file': 'Screenshot 2026-10-20 122102.png',
            'title': 'Pabellón de Gastronomía Tradicional (Distribución en Perspectiva)',
            'desc': 'Perspectiva angular de Gastronomía mostrando la distribución de mesas, letreros y vajilla.'
        },
        {
            'num': 5,
            'file': 'Screenshot 2026-10-20 125827.png',
            'title': 'Pabellón de Bienvenida (Kalihuey)',
            'desc': 'Vista frontal de Bienvenida (Kalihuey) con letrero superior, pebeteros y tapete circular.'
        },
        {
            'num': 6,
            'file': 'Screenshot 2026-10-20 125832.png',
            'title': 'Pabellón de Gastronomía Tradicional (Vista Frontal Completa)',
            'desc': 'Vista frontal de Gastronomía con letrero colgante, mesas de platillos y repisa central.'
        },
        {
            'num': 7,
            'file': 'Screenshot 2026-10-20 125835.png',
            'title': 'Transición Panteón de Deidades - Pabellón de Gastronomía',
            'desc': 'Vista de transición angular entre el Panteón de Deidades y las mesas de gastronomía.'
        },
        {
            'num': 8,
            'file': 'Screenshot 2026-10-20 125840.png',
            'title': 'Panorámica Angular: Deidades Derecha y Gastronomía',
            'desc': 'Vista panorámica angular abarcando las deidades Kauyumari, Tayau, Tatewari y gastronomía.'
        },
        {
            'num': 9,
            'file': 'Screenshot 2026-10-20 125845.png',
            'title': 'Panteón de Deidades Sagradas (Vista Frontal Completa)',
            'desc': 'Vista frontal completa del Panteón con las seis Deidades Sagradas de la cosmovisión.'
        },
        {
            'num': 10,
            'file': 'Screenshot 2026-10-20 125852.png',
            'title': 'Transición Pabellón de Fiestas - Panteón de Deidades',
            'desc': 'Vista angular de transición espacial entre el Pabellón de Fiestas y el Panteón de Deidades.'
        },
        {
            'num': 11,
            'file': 'Screenshot 2026-10-20 125857.png',
            'title': 'Pabellón de Fiestas y Ceremonias (Vista Frontal Completa)',
            'desc': 'Vista frontal completa de Fiestas y Ceremonias con sus tres altares y tambor central.'
        },
        {
            'num': 12,
            'file': 'Screenshot 2026-10-20 125902.png',
            'title': 'Transición Pabellón de Bienvenida - Pabellón de Fiestas',
            'desc': 'Vista angular de transición espacial entre el Pabellón de Bienvenida y el área de Fiestas.'
        },
        {
            'num': 13,
            'file': 'Screenshot 2026-10-20 125908.png',
            'title': 'Tapete Ceremonial Circular (Vista Cenital Superior)',
            'desc': 'Vista cenital superior del suelo del museo con el tapete ceremonial circular Wixárika.'
        },
        {
            'num': 14,
            'file': 'Screenshot 2026-10-20 125916.png',
            'title': 'Domo y Óculo Sagrado de Tatewari (Vista Cenital Inferior)',
            'desc': 'Vista cenital hacia el techo cónico de vigas de madera con el óculo sagrado de Tatewari.'
        }
    ]

    # --- PÁGINA 1: PORTADA Y TABLA RESUMEN ---
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(12)
    p_title.paragraph_format.space_after = Pt(4)
    run_title = p_title.add_run('CATÁLOGO DE VISTAS DEL MODELO 3D: MUSEO VIRTUAL WIXÁRIKA')
    run_title.bold = True
    run_title.font.size = Pt(16)
    run_title.font.color.rgb = black

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(14)
    run_sub = p_sub.add_run('Descripción visual y espacial de las 14 perspectivas capturadas en el modelo 3D')
    run_sub.font.size = Pt(11)
    run_sub.font.color.rgb = black

    # Summary table
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = table.rows[0].cells
    headers = ['Vista', 'Archivo de Captura', 'Área / Perspectiva', 'Descripción Breve (1 renglón)']
    col_widths = [Inches(0.6), Inches(1.8), Inches(1.9), Inches(2.7)]

    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        hdr_cells[i].paragraphs[0].runs[0].bold = True
        hdr_cells[i].paragraphs[0].runs[0].font.size = Pt(9.5)
        hdr_cells[i].paragraphs[0].runs[0].font.color.rgb = black
        set_cell_margins(hdr_cells[i], top=80, bottom=80, left=100, right=100)

    for v in views_data:
        row_cells = table.add_row().cells
        row_cells[0].text = f"#{v['num']}"
        row_cells[1].text = v['file']
        row_cells[2].text = v['title']
        row_cells[3].text = v['desc']
        for i in range(4):
            p = row_cells[i].paragraphs[0]
            if len(p.runs) > 0:
                p.runs[0].font.size = Pt(9)
                p.runs[0].font.color.rgb = black
                if i == 0 or i == 2:
                    p.runs[0].bold = True
            set_cell_margins(row_cells[i], top=50, bottom=50, left=100, right=100)

    # Set column widths
    for row in table.rows:
        for i, w in enumerate(col_widths):
            row.cells[i].width = w

    doc.add_page_break()

    # --- PÁGINAS INDIVIDUALES PARA CADA VISTA ---
    img_dir = os.path.join('imgsParaArchivo')
    for idx, v in enumerate(views_data):
        img_path = os.path.join(img_dir, v['file'])
        
        # 1. Título de la vista
        p_head = doc.add_paragraph()
        p_head.paragraph_format.space_before = Pt(8)
        p_head.paragraph_format.space_after = Pt(2)
        r_head = p_head.add_run(f"Vista {v['num']}: {v['title']}")
        r_head.bold = True
        r_head.font.size = Pt(13)
        r_head.font.color.rgb = black
        
        # 2. Referencia de archivo
        p_file = doc.add_paragraph()
        p_file.paragraph_format.space_before = Pt(0)
        p_file.paragraph_format.space_after = Pt(8)
        r_file_lbl = p_file.add_run("Archivo: ")
        r_file_lbl.bold = True
        r_file_lbl.font.size = Pt(9.5)
        r_file_lbl.font.color.rgb = black
        r_file_val = p_file.add_run(v['file'])
        r_file_val.font.size = Pt(9.5)
        r_file_val.font.color.rgb = black

        # 3. Imagen
        if os.path.exists(img_path):
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.paragraph_format.space_before = Pt(0)
            p_img.paragraph_format.space_after = Pt(8)
            run_img = p_img.add_run()
            run_img.add_picture(img_path, width=Inches(5.8))
        else:
            p_missing = doc.add_paragraph()
            r_miss = p_missing.add_run(f"[Imagen no encontrada: {img_path}]")
            r_miss.font.color.rgb = black

        # 4. Descripción breve (1 renglón máximo)
        p_desc = doc.add_paragraph()
        p_desc.paragraph_format.space_before = Pt(4)
        p_desc.paragraph_format.space_after = Pt(12)
        r_lbl = p_desc.add_run("Descripción: ")
        r_lbl.bold = True
        r_lbl.font.size = Pt(11)
        r_lbl.font.color.rgb = black
        r_desc = p_desc.add_run(v['desc'])
        r_desc.font.size = Pt(11)
        r_desc.font.color.rgb = black

        if idx < len(views_data) - 1:
            doc.add_page_break()

    out_file1 = 'Vistas_Modelo_3D_Museo_Wixarika.docx'
    out_file2 = os.path.join('imgsParaArchivo', 'Vistas_Modelo_3D_Museo_Wixarika.docx')
    doc.save(out_file1)
    doc.save(out_file2)
    print(f"Documento guardado con éxito en: {out_file1} y {out_file2}")

if __name__ == '__main__':
    create_document()
